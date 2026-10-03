import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from pptx import Presentation

from app.citations import verify_citations
from app.config import get_settings, ensure_structure
from app.db import Database
from app.extractors import extract_sections
from app.retrieval import index_one, inspect_search
from app.security import validate_file_content


class CitationVerifierTests(unittest.TestCase):
    def test_supported_reference_scores_high(self):
        sources = [{"text": "BitLocker permite recuperar la unidad mediante una clave de recuperación."}]
        report = verify_citations("La unidad puede recuperarse con una clave BitLocker [S1].", sources, 0.05)
        self.assertEqual(report["invalid_references"], [])
        self.assertEqual(report["unsupported_claims"], 0)
        self.assertIn(report["confidence"], {"medium", "high"})

    def test_invalid_reference_is_detected(self):
        report = verify_citations("La respuesta depende de una fuente [S9].", [{"text": "otra cosa"}], 0.05)
        self.assertEqual(report["invalid_references"], [9])
        self.assertEqual(report["confidence"], "low")


class FormatAwareExtractionTests(unittest.TestCase):
    def test_xlsx_and_pptx_are_validated_and_extracted(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            xlsx = root / "inventario.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Stock"
            ws.append(["Equipo", "Estado"])
            ws.append(["Portátil", "Disponible"])
            wb.save(xlsx)
            self.assertTrue(validate_file_content(xlsx))
            xsections = extract_sections(xlsx)
            self.assertTrue(xsections)
            self.assertIn("Stock", xsections[0]["locator"])

            pptx = root / "runbook.pptx"
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = "VPN"
            slide.placeholders[1].text = "Revisar credenciales y conectividad"
            prs.save(pptx)
            self.assertTrue(validate_file_content(pptx))
            psections = extract_sections(pptx)
            self.assertTrue(psections)
            self.assertIn("diapositiva 1", psections[0]["locator"])

    def test_json_html_csv_and_eml_extract(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            samples = {
                "data.json": json.dumps({"VPN": {"paso": "limpiar credenciales"}}, ensure_ascii=False),
                "page.html": "<html><body><h1>VPN</h1><p>Revisar ADFS.</p></body></html>",
                "table.csv": "clave,valor\nVPN,activa\n",
                "mail.eml": "From: a@example.com\nTo: b@example.com\nSubject: VPN\n\nRevisar credenciales.",
            }
            for name, content in samples.items():
                p = root / name
                p.write_text(content, encoding="utf-8")
                self.assertTrue(validate_file_content(p))
                self.assertTrue(extract_sections(p), name)


class RAGTraceTests(unittest.IsolatedAsyncioTestCase):
    async def test_inspector_exposes_pipeline_trace(self):
        with tempfile.TemporaryDirectory() as td, patch.dict(os.environ, {
            "BRAIN_ROOT": td,
            "EMBEDDING_PROVIDER": "none",
            "AUTO_INDEX": "false",
        }, clear=False):
            s = get_settings()
            ensure_structure(s)
            db = Database(Path(td) / "90_SYSTEM/database/test.db")
            p = Path(td) / "10_WIKI" / "VPN.md"
            p.write_text("# VPN\nRevisar credenciales cacheadas y conectividad ADFS.", encoding="utf-8")
            await index_one(db, s, p)
            result = await inspect_search(db, s, "credenciales VPN", 3)
            self.assertTrue(result["items"])
            trace = result["trace"]
            self.assertGreaterEqual(trace["lexical_candidates"], 1)
            self.assertIn("timings_ms", trace)
            self.assertIn("total", trace["timings_ms"])


if __name__ == "__main__":
    unittest.main()
